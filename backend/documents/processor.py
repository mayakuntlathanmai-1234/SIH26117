import os
from pypdf import PdfReader
import pytesseract
from pdf2image import convert_from_path
import tempfile

def extract_text_from_pdf(pdf_path: str) -> list:
    """
    Extracts text from a PDF. If no text is found, uses OCR.
    Returns a list of dicts with page text and metadata.
    """
    if not os.path.exists(pdf_path):
        return []
    
    extracted_data = []
    has_text = False
    
    try:
        reader = PdfReader(pdf_path)
        for page_num, page in enumerate(reader.pages):
            text = page.extract_text()
            if text and len(text.strip()) > 20:
                has_text = True
                _process_text_chunk(text, page_num + 1, pdf_path, extracted_data)
                
        # If no text was extracted (e.g., scanned PDF), fall back to OCR
        if not has_text:
            print(f"No text found in {pdf_path}. Attempting OCR...")
            _extract_via_ocr(pdf_path, extracted_data)
            
    except Exception as e:
        print(f"Error reading PDF {pdf_path}: {e}")
        
    return extracted_data

def split_text_with_overlap(text: str, chunk_size: int = 600, chunk_overlap: int = 100) -> list:
    """Splits text into overlapping chunks to preserve context across section boundaries."""
    text_clean = text.strip()
    if not text_clean:
        return []
    if len(text_clean) <= chunk_size:
        return [text_clean] if len(text_clean) > 20 else []
    
    chunks = []
    start = 0
    while start < len(text_clean):
        end = min(start + chunk_size, len(text_clean))
        if end < len(text_clean):
            # Attempt to break on sentence or paragraph boundary
            break_point = max(text_clean.rfind('\n', start, end), text_clean.rfind('. ', start, end))
            if break_point > start + (chunk_size // 2):
                end = break_point + 1
        
        chunk = text_clean[start:end].strip()
        if len(chunk) > 20:
            chunks.append(chunk)
        start += (chunk_size - chunk_overlap)
        if start >= len(text_clean) - chunk_overlap:
            break
            
    return chunks

def _process_text_chunk(text: str, page_num: int, pdf_path: str, extracted_data: list):
    chunks = split_text_with_overlap(text, chunk_size=600, chunk_overlap=100)
    for idx, chunk_text in enumerate(chunks):
        extracted_data.append({
            "text": chunk_text,
            "metadata": {
                "source": os.path.basename(pdf_path),
                "page": page_num,
                "chunk": idx
            }
        })

def _extract_via_ocr(pdf_path: str, extracted_data: list):
    try:
        # Note: poppler must be installed for pdf2image to work, and tesseract for pytesseract
        with tempfile.TemporaryDirectory() as path:
            images = convert_from_path(pdf_path, output_folder=path)
            for page_num, image in enumerate(images):
                # We can also get confidence scores by using image_to_data, but for simplicity we use image_to_string
                # and add a heuristic for low confidence.
                text = pytesseract.image_to_string(image)
                if not text or len(text.strip()) < 10:
                    text = "OCR CONFIDENCE: LOW - Unable to reliably interpret this handwritten section. Manual verification required."
                _process_text_chunk(text, page_num + 1, pdf_path, extracted_data)
    except Exception as e:
        print(f"OCR failed for {pdf_path}. Make sure Tesseract and Poppler are installed. Error: {e}")
        extracted_data.append({
            "text": "OCR CONFIDENCE: LOW - OCR Engine failed or not installed. Manual verification required.",
            "metadata": {
                "source": os.path.basename(pdf_path),
                "page": 1,
                "chunk": 0
            }
        })

def extract_text_from_image(image_path: str) -> list:
    """Extracts text from an image using local OCR and returns structured chunk."""
    if not os.path.exists(image_path):
        return []
    extracted_data = []
    try:
        text = pytesseract.image_to_string(image_path)
        if text and len(text.strip()) > 5:
            _process_text_chunk(text, 1, image_path, extracted_data)
        else:
            _process_text_chunk("OCR CONFIDENCE: LOW - Minimal text detected in image.", 1, image_path, extracted_data)
    except Exception as e:
        print(f"OCR failed for image {image_path}: {e}")
        _process_text_chunk("OCR CONFIDENCE: LOW - OCR Engine failed or not installed.", 1, image_path, extracted_data)
    return extracted_data

def extract_text_from_docx(docx_path: str) -> list:
    """Extracts paragraphs and table data from a Word (.docx) document."""
    if not os.path.exists(docx_path):
        return []
    extracted_data = []
    try:
        import docx
        doc = docx.Document(docx_path)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    full_text.append(row_text)
        
        text_content = "\n\n".join(full_text)
        if text_content:
            _process_text_chunk(text_content, 1, docx_path, extracted_data)
    except Exception as e:
        print(f"Error parsing DOCX file {docx_path}: {e}")
        # Plain text fallback
        try:
            with open(docx_path, "r", encoding="utf-8", errors="ignore") as f:
                raw = f.read()
                if len(raw.strip()) > 10:
                    _process_text_chunk(raw, 1, docx_path, extracted_data)
        except Exception:
            pass
    return extracted_data

def extract_text_from_txt(txt_path: str) -> list:
    """Extracts text from a plain text or markdown file."""
    if not os.path.exists(txt_path):
        return []
    extracted_data = []
    try:
        with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
            if content.strip():
                _process_text_chunk(content, 1, txt_path, extracted_data)
    except Exception as e:
        print(f"Error reading TXT file {txt_path}: {e}")
    return extracted_data

def process_document(file_path: str) -> list:
    """Universal file parser router based on file extension."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext in [".txt", ".md", ".csv", ".json"]:
        return extract_text_from_txt(file_path)
    elif ext in [".png", ".jpg", ".jpeg"]:
        return extract_text_from_image(file_path)
    else:
        print(f"Unsupported document extension: {ext}")
        return []

