import os
from pptx import Presentation
from pptx.util import Inches, Pt
from datetime import datetime

def create_ppt_report(data: dict, output_path: str) -> str:
    """Creates a PowerPoint presentation from structured JSON."""
    prs = Presentation()
    
    # Slide 1: Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = data.get('title', 'AI Analysis Presentation')
    subtitle.text = f"Sovereign AI Workflow\nDate: {datetime.now().strftime('%Y-%m-%d')}"
    
    bullet_slide_layout = prs.slide_layouts[1]
    
    # Generate generic slides from sections
    sections = [
        ("Executive Summary", data.get("executive_summary")),
        ("Inspection Overview", data.get("inspection_overview")),
        ("Risk Assessment", data.get("risk_assessment")),
        ("Conclusion", data.get("conclusion"))
    ]
    
    for section_title, section_content in sections:
        if section_content:
            slide = prs.slides.add_slide(bullet_slide_layout)
            slide.shapes.title.text = section_title
            tf = slide.placeholders[1].text_frame
            tf.text = str(section_content)
            
    # Generate list slides
    lists = [
        ("Key Findings", data.get("key_findings")),
        ("Technical Observations", data.get("technical_observations")),
        ("Recommended Actions", data.get("recommended_actions"))
    ]
    
    for list_title, list_items in lists:
        if list_items and isinstance(list_items, list) and len(list_items) > 0:
            slide = prs.slides.add_slide(bullet_slide_layout)
            slide.shapes.title.text = list_title
            tf = slide.placeholders[1].text_frame
            for i, item in enumerate(list_items):
                p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                p.text = str(item)
                p.level = 0
                
    prs.save(output_path)
    return output_path
