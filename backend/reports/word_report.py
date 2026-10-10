from docx import Document
import os
from datetime import datetime

def create_word_report(data: dict, output_path: str) -> str:
    doc = Document()
    doc.add_heading(data.get('title', 'AI Analysis Report'), 0)
    
    doc.add_paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    
    if data.get('subject'):
        doc.add_heading('Subject', level=1)
        doc.add_paragraph(data.get('subject'))
        
    if data.get('background'):
        doc.add_heading('Background', level=1)
        doc.add_paragraph(data.get('background'))
        
    if data.get('findings'):
        doc.add_heading('Inspection Findings', level=1)
        for finding in data.get('findings', []):
            doc.add_paragraph(f"- {finding}", style='List Bullet')
            
    if data.get('technical_assessment'):
        doc.add_heading('Technical Assessment', level=1)
        doc.add_paragraph(data.get('technical_assessment'))
        
    if data.get('risk_impact') or data.get('risk_assessment'):
        doc.add_heading('Risk/Impact', level=1)
        doc.add_paragraph(data.get('risk_impact') or data.get('risk_assessment'))
        
    if data.get('recommendations'):
        doc.add_heading('Recommended Actions', level=1)
        for rec in data.get('recommendations', []):
            doc.add_paragraph(f"- {rec}", style='List Bullet')
            
    if data.get('approval_required'):
        doc.add_heading('Required Approval', level=1)
        doc.add_paragraph(data.get('approval_required'))
        
    if data.get('conclusion'):
        doc.add_heading('Conclusion', level=1)
        doc.add_paragraph(data.get('conclusion'))
            
    doc.save(output_path)
    return output_path
